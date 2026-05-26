#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Patched preprocessing pipeline for later NHANES cycles.

Input:
D:/47/472/New-Papers/Integrated Personalized Disability-Aware/Experiments/data_processed_later_cycles

Outputs:
analysis, filled, validation, corrected folders under the same directory.

Requirements:
pip install pandas numpy scikit-learn openpyxl
"""

from __future__ import annotations

import json
import warnings
from pathlib import Path
from typing import Optional, List, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from sklearn.experimental import enable_iterative_imputer  # noqa
from sklearn.impute import IterativeImputer, KNNImputer


BASE_DIR = Path(r"D:\47\472\New-Papers\Integrated Personalized Disability-Aware\Experiments\data_processed_later_cycles")
ANALYSIS_DIR = BASE_DIR / "analysis"
FILLED_DIR = BASE_DIR / "filled"
VALIDATION_DIR = BASE_DIR / "validation"
CORRECTED_DIR = BASE_DIR / "corrected"

for folder in [ANALYSIS_DIR, FILLED_DIR, VALIDATION_DIR, CORRECTED_DIR]:
    folder.mkdir(parents=True, exist_ok=True)


NEVER_FILL = {
    "SEQN", "cycle", "target_source",
    "disability_stage",
    "upper_limb_difficulty", "mobility_difficulty", "adl_difficulty",
    "fallback_function_risk", "self_rated_health_risk",
}

CATEGORICAL_LIKE = {
    "sex", "race_ethnicity", "education", "smoking_status",
    "meets_activity_guideline", "PHQ9_category", "arthritis",
    "arthritis_type", "polypharmacy",
}


def discover_input_files() -> List[Path]:
    files = sorted(BASE_DIR.glob("NHANES_*_final_processed.csv"))
    files = [p for p in files if "pooled" not in p.name.lower()]
    if not files:
        raise FileNotFoundError(f"No cycle-level final processed CSV files found in {BASE_DIR}")
    return files


def find_rules_file() -> Optional[Path]:
    candidates = [
        Path(r"D:\47\472\New-Papers\Integrated Personalized Disability-Aware\Experiments\rules\Fill_Not_Fill.csv"),
        Path(r"D:\47\472\New-Papers\Integrated Personalized Disability-Aware\Fill_Not_Fill.csv"),
        Path(r"D:\47\472\New-Papers\Atlam1-2026\rules\Fill_Not_Fill.csv"),
        Path(r"D:\47\472\New-Papers\Atlam1-2026\Fill_Not_Fill.csv"),
        BASE_DIR / "Fill_Not_Fill.csv",
    ]
    return next((p for p in candidates if p.exists()), None)


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c).strip() for c in out.columns]
    return out


def fix_tiny_values_to_zero(series: pd.Series, threshold: float = 1e-10) -> pd.Series:
    s = pd.to_numeric(series, errors="coerce")
    return s.mask(s.abs() < threshold, 0)


def bmi_from_height_weight(weight_kg, height_cm):
    if pd.isna(weight_kg) or pd.isna(height_cm) or height_cm <= 0:
        return np.nan
    return weight_kg / ((height_cm / 100.0) ** 2)


def derive_phq9_category(score):
    if pd.isna(score):
        return np.nan
    if score < 5:
        return 0
    if score < 10:
        return 1
    if score < 15:
        return 2
    if score < 20:
        return 3
    return 4


def derive_meets_activity_guideline(met):
    if pd.isna(met):
        return np.nan
    return 1 if met >= 600 else 0


def derive_polypharmacy(count):
    if pd.isna(count):
        return np.nan
    return 1 if count >= 5 else 0


def fill_mode(series: pd.Series) -> pd.Series:
    vals = series.mode(dropna=True)
    return series.fillna(vals.iloc[0]) if not vals.empty else series


def classify_missing_degree(pct: float) -> str:
    if pct == 0:
        return "None"
    if pct < 10:
        return "Low"
    if pct < 50:
        return "Moderate"
    return "Very High"


def infer_missing_pattern(col: str, pct: float) -> str:
    c = col.lower()
    if any(x in c for x in ["income", "arthritis", "medication", "polypharmacy"]):
        return "MNAR"
    if any(x in c for x in ["ldl", "glucose", "triglycer", "cholesterol", "grip", "bmi", "hba1c", "sbp", "dbp"]):
        return "MAR"
    if any(x in c for x in ["age", "sex", "race"]):
        return "MCAR"
    return "MAR" if pct > 0 else "MCAR"


def missing_values_analysis(df: pd.DataFrame, out_path: Path) -> pd.DataFrame:
    rows = []
    n = len(df)
    for col in df.columns:
        count = int(df[col].isna().sum())
        pct = round((count / n) * 100, 2) if n else 0
        rows.append({
            "Column": col,
            "Type of Variable": "Continuous" if pd.api.types.is_numeric_dtype(df[col]) else "Categorical",
            "Likely Pattern": infer_missing_pattern(col, pct),
            "Interpretation": "Complete data" if pct == 0 else ("Highly sparse or conditional measurement" if pct > 50 else ("Moderate missingness" if pct > 10 else "Minor missingness")),
            "Missing Count": count,
            "Missing %": pct,
            "Degree of Missing": classify_missing_degree(pct),
        })
    result = pd.DataFrame(rows).sort_values("Missing %", ascending=False)
    result.to_csv(out_path, index=False, encoding="utf-8-sig")
    return result


def load_fill_rules(df: pd.DataFrame, rules_file: Optional[Path]) -> dict:
    default = {c: (0 if c in NEVER_FILL else 1) for c in df.columns}
    if rules_file is None:
        return default

    rules = pd.read_csv(rules_file)
    rules = normalize_columns(rules)
    attr = "Attribute"
    flag = "Fill (0: Do not fill/ 1: can be filled)"
    if attr not in rules.columns or flag not in rules.columns:
        return default

    raw = dict(zip(rules[attr].astype(str).str.strip(), rules[flag]))
    out = {}
    for c in df.columns:
        if c in NEVER_FILL:
            out[c] = 0
        elif c in raw:
            try:
                out[c] = int(raw[c])
            except Exception:
                out[c] = 1
        else:
            out[c] = 1
    return out


def apply_font_to_mask(source_csv: Path, output_xlsx: Path, mask: pd.DataFrame, color: str) -> None:
    df = pd.read_csv(source_csv)
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")

    wb = load_workbook(output_xlsx)
    ws = wb["Data"]
    font = Font(color=color)
    col_index = {name: idx + 1 for idx, name in enumerate(df.columns)}

    for r in range(mask.shape[0]):
        true_cols = mask.columns[mask.iloc[r].values].tolist()
        for col in true_cols:
            if col in col_index:
                ws.cell(row=r + 2, column=col_index[col]).font = font
    wb.save(output_xlsx)


def impute_dataframe(df: pd.DataFrame, rules_file: Optional[Path]) -> Tuple[pd.DataFrame, pd.DataFrame]:
    df = normalize_columns(df)
    original_missing = df.isna()
    filled = df.copy()
    rules = load_fill_rules(filled, rules_file)

    fill_cols = [c for c, flag in rules.items() if flag == 1 and c in filled.columns and filled[c].isna().any()]

    for c in fill_cols:
        if c not in CATEGORICAL_LIKE:
            filled[c] = pd.to_numeric(filled[c], errors="coerce")

    # Derived fields before imputation if possible.
    if {"PHQ9_score", "PHQ9_category"}.issubset(filled.columns):
        m = filled["PHQ9_category"].isna() & filled["PHQ9_score"].notna()
        filled.loc[m, "PHQ9_category"] = filled.loc[m, "PHQ9_score"].apply(derive_phq9_category)

    if {"MET_min_week", "meets_activity_guideline"}.issubset(filled.columns):
        m = filled["meets_activity_guideline"].isna() & filled["MET_min_week"].notna()
        filled.loc[m, "meets_activity_guideline"] = filled.loc[m, "MET_min_week"].apply(derive_meets_activity_guideline)

    if {"medication_count", "polypharmacy"}.issubset(filled.columns):
        m = filled["polypharmacy"].isna() & filled["medication_count"].notna()
        filled.loc[m, "polypharmacy"] = filled.loc[m, "medication_count"].apply(derive_polypharmacy)

    cat_cols = [c for c in fill_cols if c in CATEGORICAL_LIKE or not pd.api.types.is_numeric_dtype(filled[c])]
    for c in cat_cols:
        filled[c] = fill_mode(filled[c])

    num_cols = [c for c in fill_cols if c not in cat_cols and pd.api.types.is_numeric_dtype(filled[c])]
    if num_cols:
        try:
            if len(num_cols) >= 3 and len(filled) >= 10:
                imputer = IterativeImputer(random_state=42, max_iter=10, sample_posterior=False)
            else:
                imputer = KNNImputer(n_neighbors=5)
            filled[num_cols] = imputer.fit_transform(filled[num_cols])
        except Exception as exc:
            print(f"Warning: multivariate imputation failed. Median fallback used. Error: {exc}")
            for c in num_cols:
                filled[c] = filled[c].fillna(filled[c].median(skipna=True))

    # Derived fields after imputation.
    if {"PHQ9_score", "PHQ9_category"}.issubset(filled.columns):
        filled["PHQ9_category"] = filled["PHQ9_score"].apply(derive_phq9_category)
    if {"MET_min_week", "meets_activity_guideline"}.issubset(filled.columns):
        filled["meets_activity_guideline"] = filled["MET_min_week"].apply(derive_meets_activity_guideline)
    if {"medication_count", "polypharmacy"}.issubset(filled.columns):
        filled["polypharmacy"] = filled["medication_count"].apply(derive_polypharmacy)

    for c in ["PHQ9_score", "income_poverty_ratio", "MET_min_week", "mean_SBP", "mean_DBP", "BMI"]:
        if c in filled.columns:
            filled[c] = fix_tiny_values_to_zero(filled[c])

    return filled, original_missing & filled.notna()


def validate_dataset(df: pd.DataFrame, out_json: Path, out_txt: Path) -> dict:
    checks = {"shape": list(df.shape)}

    if {"BMI", "weight_kg", "height_cm"}.issubset(df.columns):
        calc = df["weight_kg"] / ((df["height_cm"] / 100) ** 2)
        checks["BMI_inconsistent_rows_error_gt_2"] = int(((df["BMI"] - calc).abs() > 2).sum())

    if {"LDL", "total_cholesterol", "HDL", "triglycerides"}.issubset(df.columns):
        calc = df["total_cholesterol"] - df["HDL"] - (df["triglycerides"] / 5)
        checks["LDL_inconsistent_rows_error_gt_20"] = int(((df["LDL"] - calc).abs() > 20).sum())
        checks["LDL_greater_than_total_cholesterol_rows"] = int((df["LDL"] > df["total_cholesterol"]).sum())

    if {"mean_DBP", "mean_SBP"}.issubset(df.columns):
        checks["DBP_greater_or_equal_SBP_rows"] = int((df["mean_DBP"] >= df["mean_SBP"]).sum())

    if "PHQ9_score" in df.columns:
        checks["Invalid_PHQ9_score_rows"] = int(((df["PHQ9_score"] < 0) | (df["PHQ9_score"] > 27)).sum())

    if {"MET_min_week", "meets_activity_guideline"}.issubset(df.columns):
        bad = (
            ((df["MET_min_week"] >= 600) & (df["meets_activity_guideline"] != 1)) |
            ((df["MET_min_week"] < 600) & (df["meets_activity_guideline"] != 0))
        )
        checks["MET_guideline_inconsistent_rows"] = int(bad.sum())

    if "disability_stage" in df.columns:
        checks["target_nonmissing_rows"] = int(df["disability_stage"].notna().sum())
        checks["target_counts"] = df["disability_stage"].value_counts(dropna=False).sort_index().to_dict()

    checks["missingness_percent"] = (df.isna().mean() * 100).round(2).sort_values(ascending=False).to_dict()

    out_json.write_text(json.dumps(checks, indent=2), encoding="utf-8")
    with open(out_txt, "w", encoding="utf-8") as f:
        f.write("NHANES later-cycle validation report\n")
        f.write("=" * 45 + "\n\n")
        for k, v in checks.items():
            if k != "missingness_percent":
                f.write(f"{k}: {v}\n")
        f.write("\nTop missingness percentages:\n")
        for k, v in list(checks["missingness_percent"].items())[:25]:
            f.write(f"{k}: {v}%\n")
    return checks


def correct_bmi_ldl(df: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
    out = df.copy()
    mask = pd.DataFrame(False, index=out.index, columns=out.columns)

    if {"BMI", "height_cm", "weight_kg"}.issubset(out.columns):
        calc = out.apply(lambda r: bmi_from_height_weight(r["weight_kg"], r["height_cm"]), axis=1)
        bad = (out["BMI"] - calc).abs() > 2
        for idx in out.index[bad]:
            if pd.notna(calc.loc[idx]):
                out.at[idx, "BMI"] = round(float(calc.loc[idx]), 4)
                mask.at[idx, "BMI"] = True

    if {"LDL", "total_cholesterol", "HDL", "triglycerides"}.issubset(out.columns):
        valid = out["total_cholesterol"].notna() & out["HDL"].notna() & out["triglycerides"].notna() & (out["triglycerides"] < 400)
        calc = out["total_cholesterol"] - out["HDL"] - (out["triglycerides"] / 5)
        bad = valid & (((out["LDL"] - calc).abs() > 20) | (out["LDL"] > out["total_cholesterol"]))
        for idx in out.index[bad]:
            out.at[idx, "LDL"] = round(float(calc.loc[idx]), 4)
            mask.at[idx, "LDL"] = True

    if "PHQ9_score" in out.columns:
        out.loc[(out["PHQ9_score"] < 0) | (out["PHQ9_score"] > 27), "PHQ9_score"] = np.nan
        if "PHQ9_category" in out.columns:
            out["PHQ9_category"] = out["PHQ9_score"].apply(derive_phq9_category)

    if {"MET_min_week", "meets_activity_guideline"}.issubset(out.columns):
        out["meets_activity_guideline"] = out["MET_min_week"].apply(derive_meets_activity_guideline)

    if {"medication_count", "polypharmacy"}.issubset(out.columns):
        out["polypharmacy"] = out["medication_count"].apply(derive_polypharmacy)

    return out, mask


def process_file(input_csv: Path, rules_file: Optional[Path]) -> Optional[Path]:
    stem = input_csv.stem
    print(f"\n================ PROCESSING {input_csv.name} ================")

    df = pd.read_csv(input_csv)
    df = normalize_columns(df)

    missing_path = ANALYSIS_DIR / f"{stem}_missing_values_analysis.csv"
    missing_values_analysis(df, missing_path)
    print(f"Saved missingness analysis: {missing_path}")

    filled_df, imputed_mask = impute_dataframe(df, rules_file)
    filled_csv = FILLED_DIR / f"{stem}_filled.csv"
    filled_df.to_csv(filled_csv, index=False, encoding="utf-8-sig")
    print(f"Saved filled CSV: {filled_csv}")

    try:
        apply_font_to_mask(filled_csv, FILLED_DIR / f"{stem}_filled_red.xlsx", imputed_mask, "FF0000")
    except Exception as exc:
        print(f"Warning: could not save red Excel file: {exc}")

    validate_dataset(
        filled_df,
        VALIDATION_DIR / f"{stem}_filled_validation.json",
        VALIDATION_DIR / f"{stem}_filled_validation.txt",
    )

    corrected_df, corrected_mask = correct_bmi_ldl(filled_df)
    corrected_csv = CORRECTED_DIR / f"{stem}_filled_corrected.csv"
    corrected_df.to_csv(corrected_csv, index=False, encoding="utf-8-sig")
    print(f"Saved corrected CSV: {corrected_csv}")

    try:
        apply_font_to_mask(corrected_csv, CORRECTED_DIR / f"{stem}_filled_corrected_blue.xlsx", corrected_mask, "0000FF")
    except Exception as exc:
        print(f"Warning: could not save blue Excel file: {exc}")

    validate_dataset(
        corrected_df,
        VALIDATION_DIR / f"{stem}_corrected_validation.json",
        VALIDATION_DIR / f"{stem}_corrected_validation.txt",
    )

    return corrected_csv


def make_pooled_corrected(paths: List[Path]) -> None:
    if not paths:
        print("No corrected files available for pooling.")
        return

    pooled = pd.concat([pd.read_csv(p) for p in paths], ignore_index=True, sort=False)
    pooled_path = CORRECTED_DIR / "NHANES_later_cycles_pooled_filled_corrected.csv"
    pooled.to_csv(pooled_path, index=False, encoding="utf-8-sig")

    missing_values_analysis(
        pooled,
        ANALYSIS_DIR / "NHANES_later_cycles_pooled_filled_corrected_missing_values_analysis.csv",
    )
    validate_dataset(
        pooled,
        VALIDATION_DIR / "NHANES_later_cycles_pooled_filled_corrected_validation.json",
        VALIDATION_DIR / "NHANES_later_cycles_pooled_filled_corrected_validation.txt",
    )

    print(f"\nSaved pooled corrected later-cycle file: {pooled_path}")
    print(f"Pooled corrected shape: {pooled.shape}")


def main():
    warnings.filterwarnings("ignore")

    print("Later NHANES post-processing pipeline")
    print(f"Input folder: {BASE_DIR}")
    print(f"Output corrected folder: {CORRECTED_DIR}")

    rules_file = find_rules_file()
    print(f"Using fill rules: {rules_file}" if rules_file else "No Fill_Not_Fill.csv found. Conservative default fill rules will be used.")

    files = discover_input_files()
    print("\nDetected cycle-level files:")
    for f in files:
        print(f"  - {f.name}")

    corrected_paths = []
    log_rows = []

    for f in files:
        try:
            out = process_file(f, rules_file)
            if out is not None:
                corrected_paths.append(out)
            log_rows.append({"input_file": str(f), "status": "success", "error": ""})
        except Exception as exc:
            print(f"FAILED: {f.name}: {exc}")
            log_rows.append({"input_file": str(f), "status": "failed", "error": str(exc)})

    log_path = BASE_DIR / "later_cycles_postprocessing_log.csv"
    pd.DataFrame(log_rows).to_csv(log_path, index=False, encoding="utf-8-sig")
    print(f"\nSaved postprocessing log: {log_path}")

    make_pooled_corrected(corrected_paths)

    print("\nDone.")
    print("Main final file:")
    print(CORRECTED_DIR / "NHANES_later_cycles_pooled_filled_corrected.csv")


if __name__ == "__main__":
    main()
